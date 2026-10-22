"""Génère backend/Historique.gs à partir d'une copie de l'ancien fichier Google Sheet (export .xlsx).

Usage : python3 tools/generer_historique.py "NAVETTE_LEO.xlsx" backend/Historique.gs

Règles :
- ARCH-COMPTE : lignes à 8 valeurs = 4 créneaux matin + 4 soir ; lignes à plus de 8 valeurs
  (à partir de mars 2025) = 8 colonnes matin + colonnes suivantes soir.
- COMPTE : colonnes B-I = matin, J-Q = soir.
- Seuls les nombres sont pris en compte (formules / commentaires ignorés).
- Plusieurs lignes le même jour : on garde le maximum par demi-journée.
- Jours à 0 le matin ET le soir ignorés (jours fermés).
- Les week-ends sont ignorés.
"""
import sys, datetime
import openpyxl

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=False)
totaux = {}

def nombres(row, debut, fin):
    return [v for v in row[debut:fin] if isinstance(v, (int, float)) and not isinstance(v, bool)]

def noter(d, demi, val):
    if val is None: return
    k = (d, demi)
    totaux[k] = max(totaux.get(k, 0), val)

for nom in ['ARCH-COMPTE', 'COMPTE']:
    if nom not in wb.sheetnames: continue
    for row in wb[nom].iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], datetime.datetime): continue
        d = row[0].date()
        if d.weekday() >= 5: continue
        vals = nombres(row, 1, 17)
        if nom == 'ARCH-COMPTE' and len(vals) <= 8:
            m = sum(nombres(row, 1, 5)) if nombres(row, 1, 5) else None
            s = sum(nombres(row, 5, 9)) if nombres(row, 5, 9) else None
        else:
            mv, sv = nombres(row, 1, 9), nombres(row, 9, 17)
            m = sum(mv) if mv else None
            s = sum(sv) if sv else None
        noter(d, 'MATIN', m)
        noter(d, 'SOIR', s)

jours = sorted({d for d, _ in totaux})
lignes = []
for d in jours:
    m, s = totaux.get((d, 'MATIN')), totaux.get((d, 'SOIR'))
    if (m or 0) == 0 and (s or 0) == 0: continue
    for demi, v in (('MATIN', m), ('SOIR', s)):
        if v is not None:
            lignes.append('["%s","%s",%d]' % (d.isoformat(), demi, v))

with open(dst, 'w', encoding='utf-8') as f:
    f.write('/**\n * NAVETTE – historique importé de l\'ancien Google Sheet (COMPTE + ARCH-COMPTE).\n')
    f.write(' * Total de passagers par demi-journée. Généré par tools/generer_historique.py\n')
    f.write(' * Du %s au %s – %d lignes.\n */\n' % (jours[0], jours[-1], len(lignes)))
    f.write('var HISTORIQUE = [\n' + ',\n'.join(lignes) + '\n];\n')
print(len(lignes), jours[0], jours[-1])
